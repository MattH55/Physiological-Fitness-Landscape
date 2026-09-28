"""Build the Phase 4 training corpus from NCI-ALMANAC and DrugComb.

DrugComb is now on disk (data/drugcomb/summary_v_1_5.csv, real download
from Zenodo record 15235991, 2026-09-27). The two real screens are
merged: comboFM's CombALMANAC_555300.csv (the original source) plus
DrugComb's own summary table, restricted to the studies within DrugComb
that are NOT NCI-ALMANAC (DrugComb's own `study_name == "ALMANAC"` rows
are the same underlying experiments as comboFM's file -- dropped here to
avoid double-counting the same real measurements under two paths).

Two real data-quality findings from inspecting DrugComb directly, not
assumed: (1) `synergy_bliss` is the literal string "0" for entire studies
(GDSC1, CTRPV2, CCLE, GCSI, GRAY, FIMM, UHNBREAST, BEATAML) -- these
studies never had synergy computed in DrugComb's own pipeline, so their
"0" is an absence marker, not a measured null-synergy result; treating it
as a real label would inject fabricated zero-synergy rows, so these
studies are excluded entirely. (2) DrugComb's `synergy_bliss` is on a
percentage-inhibition scale (ALMANAC-within-DrugComb ranges roughly
-38..+34, mean ~1.1) while the existing ALMANAC-via-comboFM corpus is
Bliss excess of *survival fraction* (a 0-1 scale, mean 0.002) -- dividing
DrugComb's value by 100 converts percentage-inhibition excess to the same
fraction scale. Sign convention matches: positive is more killing/
inhibition than Bliss independence in both sources.

Pairs are kept when at least one drug is a usable mimetic from
data/lincs/mimetics.json, both drugs have a Phase 1 MCF7 signature, and
the cell line is in the NCI-60 RNA-seq table. For DrugComb this real
filter is the actual bottleneck, not disk access: only 4 of the 33
registered modifier signatures (Phase 3) ever resolved to a usable
mimetic at all, so "more raw combination data" cannot multiply the
corpus past that structural ceiling -- it can only deepen coverage for
those same ~4 anchor compounds. That is reported honestly below rather
than treated as a bug.

Schedule is simultaneous and the interval is 0 for every row from both
sources (DrugComb's summary table does not record schedule/timing either).
Mutation profiles for these NCI-60 lines are not in the downloaded table,
so that feature is null.
"""
from __future__ import annotations

import csv
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone

import numpy as np
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from synlethality import config
from synlethality.featurize import featurize
from synlethality.signature_space import CANONICAL_SYMBOLS, normalize_z, to_canonical

ALMANAC = os.path.join(config.DATA_DIR, "combofm", "CombALMANAC_555300.csv")
DRUGCOMB = os.path.join(config.DATA_DIR, "drugcomb", "summary_v_1_5.csv")
MIMETICS = os.path.join(config.DATA_DIR, "lincs", "mimetics.json")
#: Real, verified 2026-09-27 by sampling synergy_bliss directly: these
#: DrugComb studies carry the literal string "0" for every row, meaning
#: synergy was never computed for them (not a real measured null result).
DRUGCOMB_ZERO_ARTIFACT_STUDIES = {
    "GDSC1", "CTRPV2", "CCLE", "GCSI", "GRAY", "FIMM", "UHNBREAST", "BEATAML",
}
#: Excluded so the same real ALMANAC measurements are not counted twice
#: (once via comboFM's file, once via DrugComb's own copy of the same study).
DRUGCOMB_EXCLUDE_STUDIES = DRUGCOMB_ZERO_ARTIFACT_STUDIES | {"ALMANAC"}
REFERENCE = os.environ.get(
    "PHASE1_REFERENCE",
    # Regenerated 2026-09-27 into the repo (see run_phase2_diagnostics.py's
    # docstring) after the external drive holding the original copy became
    # unavailable between sessions -- a personal-drive default is exactly
    # the kind of reproducibility fragility this project's provenance rule
    # exists to avoid.
    os.path.join(config.DATA_DIR, "lincs_phase1_raw", "phase1_mcf7_10uM_24h.json"),
)
EXPR = os.path.join(config.DATA_DIR, "nci60", "nci60_rnaseq_composite.xlsx")
OUT_DIR = os.path.join(config.DATA_DIR, "lincs", "phase4")
REPORT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "PHASE_4_REPORT.md",
)
SALT = re.compile(
    r"\b(hydrochloride|sulfate|sulphate|citrate|disodium|tosylate|"
    r"tartrate|phosphate sodium|sodium)\b"
)


def norm_name(value: str) -> str:
    text = SALT.sub(" ", value.lower().replace("-", " ").replace("_", " "))
    return re.sub(r"[^a-z0-9]+", "", text)


def viability(growth: float) -> float:
    return max(0.0, min(1.0, (growth + 100.0) / 200.0))


def load_baselines() -> dict[str, list[float | None]]:
    wb = openpyxl.load_workbook(EXPR, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    for _ in range(9):
        next(rows)
    header = list(next(rows))
    # columns 2.. are cell lines; column 0 is the gene symbol
    cells = []
    for col, name in enumerate(header):
        if col >= 2 and isinstance(name, str) and name.strip():
            cells.append((col, name.strip()))
    by_cell = {name: {} for _, name in cells}
    for row in rows:
        if not row or not isinstance(row[0], str):
            continue
        symbol = row[0].strip()
        if not symbol or symbol.endswith("-AS1") or symbol.endswith("-AS"):
            continue
        for col, name in cells:
            value = row[col] if col < len(row) else None
            if isinstance(value, (int, float)):
                by_cell[name][symbol] = float(np.log2(value + 1.0))
    wb.close()
    out = {}
    for name, raw in by_cell.items():
        projected = to_canonical(raw)
        if projected["coverage"] < 0.7:
            continue
        out[norm_name(name)] = {
            "name": name,
            "vector": normalize_z(projected["vector"]),
            "coverage": projected["coverage"],
        }
    return out


def load_drug_signatures() -> dict[str, list[float | None]]:
    reference = json.load(open(REFERENCE, encoding="utf-8"))
    out = {}
    for name, rec in reference.items():
        projected = to_canonical(rec["gene_zscore"])
        if projected["coverage"] < 0.7:
            continue
        out[norm_name(name)] = normalize_z(projected["vector"])
    return out


def load_mimetic_names() -> set[str]:
    payload = json.load(open(MIMETICS, encoding="utf-8"))
    names = set()
    for mod in payload["modifiers"]:
        for mim in mod["mimetics"]:
            if mim["usable"]:
                names.add(norm_name(mim["compound"]))
    return names


def bliss_excess(rows: list[dict]) -> float | None:
    mono_a = {}
    mono_b = {}
    combos = []
    for row in rows:
        c1 = float(row["Conc1"])
        c2 = float(row["Conc2"])
        growth = float(row["PercentageGrowth"])
        if c1 > 0 and c2 == 0:
            mono_a[c1] = viability(growth)
        elif c2 > 0 and c1 == 0:
            mono_b[c2] = viability(growth)
        elif c1 > 0 and c2 > 0:
            combos.append((c1, c2, viability(growth)))
    excesses = []
    for c1, c2, observed in combos:
        if c1 not in mono_a or c2 not in mono_b:
            continue
        expected = mono_a[c1] * mono_b[c2]
        excesses.append(expected - observed)
    if not excesses:
        return None
    return float(sum(excesses) / len(excesses))


def load_drugcomb_grouped() -> dict[tuple[str, str, str], list[float]]:
    """Real DrugComb synergy_bliss values, grouped by (drug_a, drug_b,
    cell_line), converted from percentage-inhibition scale to the same
    fraction-of-survival scale as the ALMANAC-via-comboFM corpus (divide
    by 100). Excludes DRUGCOMB_EXCLUDE_STUDIES (zero-artifact studies and
    DrugComb's own copy of ALMANAC)."""
    grouped: dict[tuple[str, str, str], list[float]] = defaultdict(list)
    with open(DRUGCOMB, encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["study_name"] in DRUGCOMB_EXCLUDE_STUDIES:
                continue
            raw = row.get("synergy_bliss")
            if not raw or raw in ("NA", "NULL", ""):
                continue
            try:
                value = float(raw)
            except ValueError:
                continue
            a, b = sorted((row["drug_row"].strip(), row["drug_col"].strip()))
            key = (a, b, row["cell_line_name"].strip())
            grouped[key].append(value / 100.0)
    return grouped


def main() -> int:
    for path in (ALMANAC, DRUGCOMB, MIMETICS, REFERENCE, EXPR):
        if not os.path.isfile(path):
            print("missing", path)
            return 1
    print("Loading NCI-60 baselines")
    baselines = load_baselines()
    print(f"  {len(baselines)} cell lines")
    print("Loading drug signatures")
    signatures = load_drug_signatures()
    mimetics = load_mimetic_names()
    print(f"  {len(signatures)} signatures, {len(mimetics)} mimetic names")

    almanac_grouped: dict[tuple[str, str, str], list] = defaultdict(list)
    with open(ALMANAC, encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            a, b = sorted((row["Drug1"].strip(), row["Drug2"].strip()))
            almanac_grouped[(a, b, row["CellLine"].strip())].append(row)

    print("Loading DrugComb (excluding zero-artifact studies and its own ALMANAC copy)")
    drugcomb_grouped = load_drugcomb_grouped()
    print(f"  {len(drugcomb_grouped)} distinct (drug, drug, cell) triples")

    kept = []
    skipped = {
        "almanac": {"no_signature": 0, "no_baseline": 0, "not_mimetic": 0, "no_score": 0},
        "drugcomb": {"no_signature": 0, "no_baseline": 0, "not_mimetic": 0, "no_score": 0},
    }

    def _try_keep(source: str, drug_a: str, drug_b: str, cell: str, score: float | None) -> None:
        na, nb, nc = norm_name(drug_a), norm_name(drug_b), norm_name(cell)
        if na not in mimetics and nb not in mimetics:
            skipped[source]["not_mimetic"] += 1
            return
        if na not in signatures or nb not in signatures:
            skipped[source]["no_signature"] += 1
            return
        if nc not in baselines:
            skipped[source]["no_baseline"] += 1
            return
        if score is None:
            skipped[source]["no_score"] += 1
            return
        features = featurize(
            signatures[na], signatures[nb], baselines[nc]["vector"],
            None, "simultaneous", 0.0,
        )
        kept.append({
            "drug_a": drug_a,
            "drug_b": drug_b,
            "cell_line": baselines[nc]["name"],
            "score": score,
            "features": features,
            "source": source,
        })

    for (drug_a, drug_b, cell), rows in almanac_grouped.items():
        _try_keep("almanac", drug_a, drug_b, cell, bliss_excess(rows))
    for (drug_a, drug_b, cell), values in drugcomb_grouped.items():
        _try_keep("drugcomb", drug_a, drug_b, cell,
                   float(sum(values) / len(values)) if values else None)

    n_almanac = sum(1 for row in kept if row["source"] == "almanac")
    n_drugcomb = sum(1 for row in kept if row["source"] == "drugcomb")
    cells = sorted({row["cell_line"] for row in kept})
    scores = np.array([row["score"] for row in kept], dtype=np.float32)
    n = len(kept)
    n_pos = int((scores > 0).sum()) if n else 0
    n_neg = int((scores < 0).sum()) if n else 0
    median = float(np.median(scores)) if n else float("nan")
    mean = float(scores.mean()) if n else float("nan")
    # Positive skew: mean above median and more than 60% of scores positive.
    positive_skew = bool(n and mean > median and n_pos / n > 0.6)
    gate = n >= 20000 and len(cells) >= 10 and not positive_skew

    os.makedirs(OUT_DIR, exist_ok=True)
    if n:
        stack = lambda key: np.array(
            [[np.nan if v is None else v for v in row["features"][key]] for row in kept],
            dtype=np.float32,
        )
        np.savez_compressed(
            os.path.join(OUT_DIR, "features.npz"),
            signature_a=stack("signature_a"),
            signature_b=stack("signature_b"),
            baseline=stack("baseline"),
            score=scores,
        )
        with open(os.path.join(OUT_DIR, "pairs.jsonl"), "w", encoding="utf-8") as fh:
            for row in kept:
                fh.write(json.dumps({
                    "drug_a": row["drug_a"],
                    "drug_b": row["drug_b"],
                    "cell_line": row["cell_line"],
                    "score": row["score"],
                    "source": row["source"],
                    "schedule": "simultaneous",
                    "interval_min": 0,
                    "mutations": None,
                }) + "\n")

    summary = {
        "built": datetime.now(timezone.utc).isoformat(),
        "screens": [ALMANAC, DRUGCOMB],
        "drugcomb_excluded_studies": sorted(DRUGCOMB_EXCLUDE_STUDIES),
        "drugcomb_zero_artifact_studies": sorted(DRUGCOMB_ZERO_ARTIFACT_STUDIES),
        "expression": EXPR,
        "expression_citation": "Reinhold et al. 2019, figshare 10.1158/0008-5472.22420701, FPKM",
        "n_pairs": n,
        "n_pairs_almanac": n_almanac,
        "n_pairs_drugcomb": n_drugcomb,
        "n_cell_lines": len(cells),
        "cell_lines": cells,
        "score_mean": mean,
        "score_median": median,
        "n_positive": n_pos,
        "n_negative": n_neg,
        "positive_skew": positive_skew,
        "skipped": skipped,
        "mutations": "null; the NCI-60 RNA-seq table has no mutation profiles",
        "schedule": "simultaneous",
        "interval_min": 0,
        "gate_4": {
            "passed": gate,
            "criterion": ">=20000 pairs, >=10 cell lines, synergy scores not positive-skewed",
        },
    }
    with open(os.path.join(OUT_DIR, "summary.json"), "w", encoding="utf-8") as fh:
        json.dump(summary, fh, indent=2)
    _report(summary)
    print(f"GATE 4: {'PASSED' if gate else 'FAILED'} pairs={n} "
          f"(almanac={n_almanac} drugcomb={n_drugcomb}) cells={len(cells)} "
          f"pos={n_pos} neg={n_neg}")
    print("skipped", skipped)
    return 0 if gate else 2


def _report(summary: dict) -> None:
    gate = summary["gate_4"]["passed"]
    lines = [
        "# Phase 4 report — training corpus",
        "",
        f"Built {summary['built']}.",
        "",
        "DrugComb (2026-09-27, Zenodo record 15235991) is now on disk alongside "
        "NCI-ALMANAC (`CombALMANAC_555300.csv`). The corpus is the mimetic-"
        "anchored subset of both: a pair is kept when at least one drug is a "
        "usable mimetic, both drugs have a Phase 1 MCF7 signature, and the cell "
        "line is in the NCI-60 RNA-seq table (Reinhold et al. 2019, figshare "
        "10.1158/0008-5472.22420701). Baselines are log2(FPKM + 1), projected "
        "onto the 978 landmark genes and robust-z scored.",
        "",
        "DrugComb's own `ALMANAC` study rows are excluded (same underlying "
        "measurements as the comboFM file, dropped here to avoid double-"
        "counting). These DrugComb studies are excluded outright: "
        f"{', '.join(summary['drugcomb_zero_artifact_studies'])} -- their "
        "`synergy_bliss` column is the literal string \"0\" for every row, "
        "meaning synergy was never computed for them, not that it was measured "
        "as zero.",
        "",
        "The score is mean Bliss excess of survival fraction. For ALMANAC this "
        "is computed from the measured monotherapy edges and combination wells "
        "(expected survival minus observed survival). For DrugComb it is the "
        "table's own `synergy_bliss` (percentage-inhibition scale) divided by "
        "100 to match. Positive is more killing than independence in both. "
        "Negative is antagonism. No sign filter was applied.",
        "",
        "Every pair is `simultaneous` with interval 0 -- neither source records "
        "timing. Mutation profiles are null; neither expression table contains "
        "them, and none were filled in.",
        "",
        f"**Gate 4: {'PASSED' if gate else 'FAILED'}.** "
        f"{summary['n_pairs']} pairs ({summary['n_pairs_almanac']} ALMANAC + "
        f"{summary['n_pairs_drugcomb']} DrugComb) across {summary['n_cell_lines']} "
        f"cell lines. Mean score {summary['score_mean']:.4f}, median "
        f"{summary['score_median']:.4f}, {summary['n_positive']} positive and "
        f"{summary['n_negative']} negative. Positive skew: {summary['positive_skew']}.",
        "",
        "The real bottleneck is the mimetic-anchoring filter itself, not disk "
        "access: only 4 of the 33 registered modifier signatures (Phase 3) ever "
        "resolved to a usable mimetic. Adding DrugComb deepens coverage for "
        "those same ~4 anchor compounds; it cannot multiply the corpus past "
        "that structural ceiling. That is the finding, not a rerun condition.",
        "",
        "Pairs dropped before featurization:",
        "",
    ]
    for source, counts in summary["skipped"].items():
        lines.append(f"- {source}:")
        for key, count in counts.items():
            lines.append(f"  - {key}: {count}")
    lines += ["", "Cell lines:", ""]
    for name in summary["cell_lines"]:
        lines.append(f"- {name}")
    lines.append("")
    with open(REPORT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


if __name__ == "__main__":
    sys.exit(main())
